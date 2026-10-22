"""The hub exports a profile's model as documentation, as the standalone app does.

metaseed 0.66 documents a profile as Markdown, a fields CSV and an Excel
workbook. The hub serves its own explorer and builder, so without routes and
buttons of its own the export existed only in the standalone app.
"""

from __future__ import annotations

import io
from pathlib import Path
from unittest.mock import AsyncMock, patch

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from starlette.routing import Mount

from metaseed_hub.auth import TokenUser
from metaseed_hub.database import get_session
from metaseed_hub.main import create_app
from metaseed_hub.ui.dependencies import get_current_user_from_cookie

TEMPLATES = Path(__file__).resolve().parents[1] / "src" / "metaseed_hub" / "ui" / "templates"
_TOKEN = TokenUser(sub="kc-doc", email="d@example.org", name="D", roles=[], entitlements=[])


async def _get(session, *paths: str) -> list:
    app = create_app()
    hub = next(r.app for r in app.routes if isinstance(r, Mount) and r.path == "/hub")
    hub.dependency_overrides[get_current_user_from_cookie] = lambda: _TOKEN

    async def _session():
        yield session

    hub.dependency_overrides[get_session] = _session
    with patch(
        "metaseed_hub.ui.dependencies.get_current_user_from_cookie",
        AsyncMock(return_value=_TOKEN),
    ):
        async with AsyncClient(transport=ASGITransport(app=app), base_url="https://test") as c:
            return [await c.get(path) for path in paths]


@pytest.mark.asyncio
async def test_a_profile_exports_as_markdown_csv_and_workbook(session) -> None:
    markdown, csv, xlsx = await _get(
        session,
        "/hub/explore/doc/markdown/ena/1.0",
        "/hub/explore/doc/csv/ena/1.0",
        "/hub/explore/doc/xlsx/ena/1.0",
    )

    assert markdown.status_code == 200
    assert "### Experiment" in markdown.text
    assert markdown.headers["content-disposition"] == 'attachment; filename="ena-1.0.md"'
    assert csv.text.splitlines()[0].startswith("Entity,Field,Type,Required")
    assert load_workbook(io.BytesIO(xlsx.content)).sheetnames == ["Entities", "Fields", "Rules"]


@pytest.mark.asyncio
async def test_an_unknown_profile_is_not_found(session) -> None:
    (response,) = await _get(session, "/hub/explore/doc/markdown/no-such-profile/9.9")

    assert response.status_code == 404


def test_the_explorer_and_the_builder_offer_it() -> None:
    explore = (TEMPLATES / "explore" / "index.html").read_text(encoding="utf-8")
    builder = (TEMPLATES / "spec_builder" / "base.html").read_text(encoding="utf-8")

    for fmt in ("markdown", "csv", "xlsx"):
        assert f"exportModel('{fmt}')" in explore
        assert f"/hub/spec-builder/{{{{ draft_id }}}}/export/doc/{fmt}" in builder
